"""Parser and row-validation tests without a live MySQL server."""

import csv
from io import BytesIO, StringIO
from decimal import Decimal
import unittest

from openpyxl import Workbook

from app.data_import import build_import_plan, normalize_datasets, parse_import_file, _pdf_table_rows


class ImportValidationCursor:
    def __init__(self, customers=None, products=None, aliases=None, prices=None, history=None):
        self.tables = {
            "customers": customers or [],
            "products": products or [],
            "aliases": aliases or [],
            "prices": prices or [],
            "history": history or [],
        }
        self.rows = []

    def execute(self, query, _params=()):
        if "FROM customers" in query:
            self.rows = self.tables["customers"]
        elif "FROM products" in query:
            self.rows = self.tables["products"]
        elif "FROM product_aliases" in query:
            self.rows = self.tables["aliases"]
        elif "FROM customer_product_prices" in query:
            self.rows = self.tables["prices"]
        elif "FROM imported_order_history" in query:
            self.rows = self.tables["history"]
        else:
            raise AssertionError(f"Unexpected import validation query: {query}")

    def fetchall(self):
        return self.rows


class DataImportParserTests(unittest.TestCase):
    def test_csv_customer_and_product_amounts_are_normalized_to_minor_units(self):
        customer_csv = b"customer_id,customer_name,credit_limit,credit_used,currency\nDEMO-C-1,Sample Shop,1000.00,25.00,INR\n"
        product_csv = b"sku,product_name,quantity_on_hand,unit,unit_price,currency,pack_size,pack_unit\nDEMO-P-1,Rice,10,kg,75.50,INR,5,kg\n"
        customers, customer_errors = normalize_datasets(parse_import_file("customers.csv", customer_csv))
        products, product_errors = normalize_datasets(parse_import_file("products.csv", product_csv))

        self.assertEqual(customer_errors, [])
        self.assertEqual(products[0]["records"][0]["unit_price_minor"], 7550)
        self.assertEqual(products[0]["records"][0]["pack_size"], "5")
        self.assertEqual(product_errors, [])

    def test_xlsx_workbook_reads_named_customer_and_product_sheets(self):
        workbook = Workbook()
        customers = workbook.active
        customers.title = "customers"
        customers.append(["customer_id", "customer_name", "credit_limit", "credit_used", "currency"])
        customers.append(["DEMO-C-1", "Sample Shop", 1000, 0, "INR"])
        products = workbook.create_sheet("products")
        products.append(["sku", "product_name", "quantity_on_hand", "unit", "unit_price", "currency"])
        products.append(["DEMO-P-1", "Rice", 10, "kg", 75.5, "INR"])
        stream = BytesIO()
        workbook.save(stream)

        parsed = parse_import_file("shop.xlsx", stream.getvalue())
        normalized, errors = normalize_datasets(parsed)

        self.assertEqual([item["type"] for item in normalized], ["customers", "products"])
        self.assertEqual(errors, [])

    def test_pdf_table_parser_accepts_pipe_separated_catalogue(self):
        records, errors = _pdf_table_rows("""SKU | Product Name | Quantity On Hand | Unit | Unit Price | Currency
DEMO-P-1 | Sample Rice | 10 | kg | 75.50 | INR""")

        self.assertEqual(errors, [])
        self.assertEqual(records[0]["values"]["sku"], "DEMO-P-1")
        self.assertEqual(records[0]["values"]["product_name"], "Sample Rice")

    def test_invalid_unit_and_missing_values_are_row_errors(self):
        content = b"sku,product_name,quantity_on_hand,unit,unit_price,currency\nDEMO-P-1,Rice,5,bag-of-magic,75.00,INR\n"
        normalized, errors = normalize_datasets(parse_import_file("products.csv", content))

        self.assertTrue(any(error["code"] == "UNIT_UNSUPPORTED" for error in errors))
        self.assertEqual(normalized[0]["records"][0]["unit"], None)

    def test_duplicate_ids_and_existing_records_require_explicit_update_choice(self):
        content = b"customer_id,customer_name,credit_limit,credit_used,currency\nC-1,Shop One,1000,0,INR\nC-1,Shop Two,1200,0,INR\n"
        normalized, parse_errors = normalize_datasets(parse_import_file("customers.csv", content))
        plan = build_import_plan(ImportValidationCursor(), normalized, False, parse_errors)

        self.assertTrue(any(error["code"] == "DUPLICATE_IDENTIFIER" for error in plan["errors"]))
        self.assertEqual(plan["summary"]["creates"], 0)

    def test_existing_customer_update_plan_includes_before_and_after_values(self):
        content = b"customer_id,customer_name,credit_limit,credit_used,currency\nC-1,Shop One,1200,100,INR\n"
        normalized, errors = normalize_datasets(parse_import_file("customers.csv", content))
        existing = {"id": "C-1", "name": "Shop One", "credit_limit_minor": 100000, "credit_used_minor": 0, "credit_currency": "INR"}
        blocked = build_import_plan(ImportValidationCursor(customers=[existing]), normalized, False, errors)
        allowed = build_import_plan(ImportValidationCursor(customers=[existing]), normalized, True, errors)

        self.assertIn("EXISTING_RECORD_REQUIRES_UPDATE_CONFIRMATION", [error["code"] for error in blocked["errors"]])
        change = allowed["plans"]["customers"][0]
        self.assertEqual(change["action"], "update")
        self.assertEqual(change["before"]["credit_limit_minor"], 100000)
        self.assertEqual(change["values"]["credit_limit_minor"], 120000)

    def test_inventory_and_customer_price_relationships_are_checked(self):
        inventory = b"sku,quantity_on_hand,unit\nMISSING-SKU,10,kg\n"
        prices = b"customer_id,sku,unit_price,currency\nMISSING-C,OTHER-SKU,1.00,INR\n"
        inv_data, inv_errors = normalize_datasets(parse_import_file("inventory.csv", inventory))
        price_data, price_errors = normalize_datasets(parse_import_file("prices.csv", prices))
        inv_plan = build_import_plan(ImportValidationCursor(), inv_data, False, inv_errors)
        price_plan = build_import_plan(ImportValidationCursor(), price_data, False, price_errors)

        self.assertIn("PRODUCT_NOT_FOUND", [error["code"] for error in inv_plan["errors"]])
        self.assertIn("PRODUCT_NOT_FOUND", [error["code"] for error in price_plan["errors"]])
        self.assertIn("CUSTOMER_NOT_FOUND", [error["code"] for error in price_plan["errors"]])

    def test_previous_order_rejects_customer_id_name_mismatch(self):
        customers = [{"id": "C-1", "name": "Correct Shop", "credit_currency": "INR"}]
        products = [{"id": "P-1", "sku": "RICE-1", "name": "Rice", "unit": "kg", "quantity_on_hand": Decimal("4"), "unit_price_minor": 5000, "price_currency": "INR", "pack_size": None, "pack_unit": None}]
        content = b"order_id,customer_id,customer_name,order_date,status,sku,quantity,unit,unit_price,currency\nOLD-1,C-1,Wrong Shop,2025-01-01,delivered,RICE-1,2,kg,50,INR\n"
        normalized, errors = normalize_datasets(parse_import_file("history.csv", content))

        plan = build_import_plan(ImportValidationCursor(customers=customers, products=products), normalized, False, errors)

        self.assertIn("CUSTOMER_ID_NAME_MISMATCH", [error["code"] for error in plan["errors"]])
        self.assertEqual(plan["plans"]["previous_orders"], [])

    def test_previous_order_rows_must_agree_on_status(self):
        customers = [{"id": "C-1", "name": "Shop", "credit_currency": "INR"}]
        products = [
            {"id": "P-1", "sku": "RICE-1", "name": "Rice", "unit": "kg", "quantity_on_hand": Decimal("4"), "unit_price_minor": 5000, "price_currency": "INR", "pack_size": None, "pack_unit": None},
            {"id": "P-2", "sku": "OIL-1", "name": "Oil", "unit": "l", "quantity_on_hand": Decimal("4"), "unit_price_minor": 5000, "price_currency": "INR", "pack_size": None, "pack_unit": None},
        ]
        content = b"order_id,customer_id,order_date,status,sku,quantity,unit,unit_price,currency\nOLD-1,C-1,2025-01-01,delivered,RICE-1,2,kg,50,INR\nOLD-1,C-1,2025-01-01,pending,OIL-1,1,l,50,INR\n"
        normalized, errors = normalize_datasets(parse_import_file("history.csv", content))

        plan = build_import_plan(ImportValidationCursor(customers=customers, products=products), normalized, False, errors)

        self.assertIn("HISTORICAL_ORDER_INCONSISTENT", [error["code"] for error in plan["errors"]])
        # Valid rows remain visible in preview, but any row error blocks the
        # entire batch from confirmation.
        self.assertEqual(len(plan["plans"]["previous_orders"]), 1)


if __name__ == "__main__":
    unittest.main()
