"""Staging and validation helpers for shopkeeper-controlled business imports."""

import csv
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO, StringIO
import re
import uuid

from openpyxl import load_workbook
from pypdf import PdfReader


MAX_IMPORT_ROWS = 5000
MAX_IMPORT_BYTES = 15 * 1024 * 1024
SUPPORTED_UNITS = {
    "g", "kg", "mg", "l", "ml", "each", "unit", "piece", "pieces", "pcs",
    "box", "case", "pack", "bottle", "carton", "bag", "bunch", "dozen", "tray",
}
_CURRENCY_DIGITS = {"BIF": 0, "CLP": 0, "DJF": 0, "GNF": 0, "ISK": 0, "JPY": 0, "KRW": 0, "PYG": 0, "UGX": 0, "VND": 0, "XAF": 0, "XOF": 0, "XPF": 0, "BHD": 3, "IQD": 3, "JOD": 3, "KWD": 3, "LYD": 3, "OMR": 3, "TND": 3}
_HEADER_ALIASES = {
    "customerid": "customer_id", "customer_code": "customer_id", "account_id": "customer_id",
    "customername": "customer_name", "account_name": "customer_name",
    "productname": "product_name", "item_name": "product_name", "description": "product_name",
    "productcode": "sku", "product_code": "sku", "item_code": "sku", "barcode": "sku",
    "stock": "quantity_on_hand", "quantity": "quantity_on_hand", "qty": "quantity_on_hand", "on_hand": "quantity_on_hand", "qty_on_hand": "quantity_on_hand",
    "price": "unit_price", "catalog_price": "unit_price", "selling_price": "unit_price",
    "creditlimit": "credit_limit", "creditused": "credit_used", "pricecurrency": "currency",
    "packsize": "pack_size", "packunit": "pack_unit", "orderid": "order_id", "orderdate": "order_date",
    "unitofmeasure": "unit", "uom": "unit", "alias_list": "aliases", "product_aliases": "aliases",
}
_SHEET_TYPES = {
    "customer": "customers", "customers": "customers", "customer_list": "customers", "product": "products", "products": "products", "product_list": "products",
    "catalog": "products", "catalogue": "products", "inventory": "inventory", "inventory_list": "inventory", "stock": "inventory",
    "price": "prices", "prices": "prices", "customer_prices": "prices", "previous_orders": "previous_orders",
    "order_history": "previous_orders", "historical_orders": "previous_orders",
}


class ImportFormatError(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def _header(value):
    key = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")
    compact = key.replace("_", "")
    return _HEADER_ALIASES.get(compact, _HEADER_ALIASES.get(key, key))


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return str(value).strip()


def _infer_type(headers, sheet_name="", forced_type=None):
    if forced_type:
        return forced_type
    sheet_key = re.sub(r"[^a-z0-9]+", "_", sheet_name.casefold()).strip("_")
    if sheet_key in _SHEET_TYPES:
        return _SHEET_TYPES[sheet_key]
    fields = {_header(h) for h in headers}
    if {"order_id", "sku"}.issubset(fields) and ({"customer_id", "customer_name"} & fields):
        return "previous_orders"
    if {"customer_id", "customer_name"}.issubset(fields) or {"customer_id", "credit_limit"}.issubset(fields):
        return "customers"
    if {"sku", "product_name", "unit_price", "quantity_on_hand"}.issubset(fields):
        return "products"
    if {"sku", "quantity_on_hand"}.issubset(fields):
        return "inventory"
    if {"sku", "unit_price"}.issubset(fields):
        return "prices"
    return None


def _rows_to_records(headers, rows, kind, sheet_name=""):
    normalized_headers = [_header(value) for value in headers]
    result = []
    errors = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        record = {key: _cell(values[index]) if index < len(values) else "" for index, key in enumerate(normalized_headers) if key}
        if kind == "customers" and not record.get("name"):
            record["name"] = record.get("customer_name", "")
        if kind in {"products", "inventory", "prices", "previous_orders"} and not record.get("name"):
            record["name"] = record.get("product_name", "")
        result.append({"sheet": sheet_name or "CSV", "row": row_number, "values": record})
        if len(result) > MAX_IMPORT_ROWS:
            errors.append({"sheet": sheet_name or "CSV", "row": row_number, "field": None, "code": "ROW_LIMIT_EXCEEDED", "message": f"An import can contain at most {MAX_IMPORT_ROWS} rows."})
            break
    return result, errors


def _pdf_table_rows(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    header_index = None
    delimiter = None
    headers = []
    for index, line in enumerate(lines):
        if "sku" not in line.casefold() and "product code" not in line.casefold():
            continue
        if "|" in line:
            delimiter = "|"
            headers = [piece.strip() for piece in line.split("|")]
        elif "\t" in line:
            delimiter = "\t"
            headers = [piece.strip() for piece in line.split("\t")]
        else:
            delimiter = "spaces"
            headers = re.split(r"\s{2,}", line)
        if len(headers) >= 4 and any("price" in h.casefold() for h in headers):
            header_index = index
            break
    if header_index is None:
        raise ImportFormatError("PDF_TABLE_NOT_FOUND", "Could not find a catalogue table with SKU, product, unit, and price columns.")
    rows = []
    errors = []
    for row_number, line in enumerate(lines[header_index + 1:], start=header_index + 3):
        if delimiter == "|":
            pieces = [piece.strip() for piece in line.split("|")]
        elif delimiter == "\t":
            pieces = [piece.strip() for piece in line.split("\t")]
        else:
            pieces = re.split(r"\s{2,}", line)
        if len(pieces) != len(headers):
            continue
        rows.append({"sheet": "PDF", "row": row_number, "values": {key: _cell(pieces[i]) for i, key in enumerate(map(_header, headers))}})
    if not rows:
        errors.append({"sheet": "PDF", "row": None, "field": None, "code": "PDF_TABLE_EMPTY", "message": "The PDF catalogue table contained no parseable data rows."})
    return rows, errors


def _extract_pdf(file_bytes):
    reader = PdfReader(BytesIO(file_bytes))
    if len(reader.pages) > 40:
        raise ImportFormatError("PAGE_LIMIT_EXCEEDED", "PDF imports are limited to 40 pages.")
    texts = [(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(texts).strip()
    if text:
        return text, "text"
    try:
        import pypdfium2 as pdfium
        import pytesseract
        document = pdfium.PdfDocument(file_bytes)
        text = "\n".join(pytesseract.image_to_string(document[index].render(scale=2).to_pil()) for index in range(len(document)))
    except (ImportError, OSError, RuntimeError) as exc:
        raise ImportFormatError("OCR_UNAVAILABLE", "This scanned PDF needs Tesseract OCR. Install the Tesseract binary and ensure it is on PATH, then retry.") from exc
    if not text.strip():
        raise ImportFormatError("OCR_NO_TEXT", "OCR did not detect readable catalogue text in this PDF.")
    return text, "ocr"


def parse_import_file(filename, file_bytes, dataset_type=None):
    """Extract rows from CSV/XLSX or a tabular PDF without writing business data."""
    if not file_bytes or len(file_bytes) > MAX_IMPORT_BYTES:
        raise ImportFormatError("FILE_SIZE_INVALID", f"Choose a non-empty file no larger than {MAX_IMPORT_BYTES // (1024 * 1024)} MB.")
    ext = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if ext not in {"csv", "xlsx", "pdf"}:
        raise ImportFormatError("UNSUPPORTED_FILE_TYPE", "Upload a CSV, XLSX, or PDF file.")
    datasets = []
    extraction = "native"
    try:
        if ext == "csv":
            content = file_bytes.decode("utf-8-sig")
            reader = csv.reader(StringIO(content))
            headers = next(reader, None)
            if not headers:
                raise ImportFormatError("FILE_EMPTY", "The CSV file is empty.")
            kind = _infer_type(headers, forced_type=dataset_type)
            if not kind:
                raise ImportFormatError("DATASET_TYPE_UNKNOWN", "CSV headers do not match customers, products, inventory, prices, or previous orders.")
            records, errors = _rows_to_records(headers, reader, kind)
            datasets.append({"type": kind, "records": records, "errors": errors})
        elif ext == "xlsx":
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
            if len(workbook.worksheets) > 5:
                raise ImportFormatError("SHEET_LIMIT_EXCEEDED", "An XLSX import can contain at most five dataset sheets.")
            for sheet in workbook.worksheets:
                iterator = sheet.iter_rows(values_only=True)
                headers = next(iterator, None)
                if not headers:
                    continue
                kind = _infer_type(headers, sheet.title, dataset_type)
                if not kind:
                    raise ImportFormatError("DATASET_TYPE_UNKNOWN", f"Could not identify import data in sheet '{sheet.title}'. Rename it to a supported type or provide recognized headers.")
                records, errors = _rows_to_records(headers, iterator, kind, sheet.title)
                datasets.append({"type": kind, "records": records, "errors": errors})
            workbook.close()
            if not datasets:
                raise ImportFormatError("FILE_EMPTY", "The workbook has no importable sheets.")
        else:
            text, extraction = _extract_pdf(file_bytes)
            records, errors = _pdf_table_rows(text)
            kind = dataset_type or "products"
            if kind not in {"products", "prices", "inventory"}:
                raise ImportFormatError("PDF_DATASET_UNSUPPORTED", "PDF imports support product catalogues, prices, or inventory only.")
            datasets.append({"type": kind, "records": records, "errors": errors})
    except UnicodeDecodeError as exc:
        raise ImportFormatError("CSV_ENCODING_INVALID", "CSV files must use UTF-8 encoding.") from exc
    except ImportFormatError:
        raise
    except Exception as exc:
        raise ImportFormatError("FILE_PARSE_FAILED", "The uploaded file could not be read. Verify it is a valid CSV, XLSX, or PDF.") from exc
    if sum(len(dataset["records"]) for dataset in datasets) > MAX_IMPORT_ROWS:
        raise ImportFormatError("ROW_LIMIT_EXCEEDED", f"An import can contain at most {MAX_IMPORT_ROWS} rows across all sheets.")
    if not any(dataset["records"] for dataset in datasets):
        raise ImportFormatError("FILE_EMPTY", "The file contains headers but no data rows.")
    return {"format": ext, "extraction": extraction, "datasets": datasets}


def _parse_decimal(value, row, field, errors, *, allow_zero=False):
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        amount = None
    if amount is None or not amount.is_finite() or amount < 0 or (amount == 0 and not allow_zero):
        errors.append({**row, "field": field, "code": "INVALID_NUMBER", "message": f"{field} must be a {'non-negative' if allow_zero else 'positive'} number."})
        return None
    if max(0, -amount.as_tuple().exponent) > 3 or amount >= Decimal("100000000000"):
        errors.append({**row, "field": field, "code": "NUMBER_PRECISION_INVALID", "message": f"{field} supports at most three decimal places."})
        return None
    return amount


def _unit(value, row, field, errors):
    unit = str(value or "").strip().casefold()
    if unit not in SUPPORTED_UNITS:
        errors.append({**row, "field": field, "code": "UNIT_UNSUPPORTED", "message": f"{field} must be a supported retail unit such as kg, l, each, box, or case."})
        return None
    return unit


def _decimal_text(value):
    return format(value, "f") if value is not None else None


def _money_minor(value, currency, row, field, errors):
    currency = str(currency or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{3}", currency):
        errors.append({**row, "field": "currency", "code": "CURRENCY_INVALID", "message": "Currency must be a three-letter ISO currency code."})
        return None, None
    text = str(value or "").strip()
    if "," in text and not re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", text):
        errors.append({**row, "field": field, "code": "NUMBER_FORMAT_INVALID", "message": f"{field} must use standard decimal notation; thousands separators must group three digits."})
        return None, currency
    amount = _parse_decimal(text, row, field, errors, allow_zero=True)
    if amount is None:
        return None, currency
    digits = _CURRENCY_DIGITS.get(currency, 2)
    scaled = amount * (Decimal(10) ** digits)
    if scaled != scaled.to_integral_value():
        errors.append({**row, "field": field, "code": "CURRENCY_PRECISION_INVALID", "message": f"{currency} supports {digits} decimal places in this import."})
        return None, currency
    return int(scaled), currency


def normalize_datasets(parsed):
    """Convert source cells to explicit schema values; return field-level errors."""
    normalized = []
    errors = []
    for dataset in parsed["datasets"]:
        kind = dataset["type"]
        for error in dataset["errors"]:
            errors.append(error)
        records = []
        for source in dataset["records"]:
            row = {"sheet": source["sheet"], "row": source["row"]}
            values = source["values"]
            record = {"type": kind, "sheet": row["sheet"], "row": row["row"]}
            if kind == "customers":
                record.update(id=values.get("customer_id") or values.get("id", ""), name=values.get("customer_name") or values.get("name", ""))
                required = ("id", "name", "credit_limit", "credit_used", "currency")
                for field in required:
                    if not str(record.get(field, values.get(field, ""))).strip():
                        errors.append({**row, "field": field, "code": "REQUIRED_FIELD_MISSING", "message": f"{field} is required."})
                for field in ("credit_limit", "credit_used"):
                    amount, currency = _money_minor(values.get(field), values.get("currency"), row, field, errors)
                    record[f"{field}_minor"] = amount
                    record["currency"] = currency
                record["id"] = str(record["id"]).strip()
                record["name"] = str(record["name"]).strip()
                for field, maximum in (("id", 64), ("name", 255)):
                    if len(record[field]) > maximum:
                        errors.append({**row, "field": field, "code": "FIELD_TOO_LONG", "message": f"{field} may contain at most {maximum} characters."})
            elif kind == "products":
                record.update(sku=str(values.get("sku", "")).strip().upper(), name=str(values.get("product_name") or values.get("name") or "").strip())
                for field in ("sku", "name", "unit", "quantity_on_hand", "unit_price", "currency"):
                    if not str(values.get(field, record.get(field, ""))).strip():
                        errors.append({**row, "field": field, "code": "REQUIRED_FIELD_MISSING", "message": f"{field} is required."})
                record["unit"] = _unit(values.get("unit"), row, "unit", errors)
                record["quantity_on_hand"] = _decimal_text(_parse_decimal(values.get("quantity_on_hand"), row, "quantity_on_hand", errors, allow_zero=True))
                record["unit_price_minor"], record["currency"] = _money_minor(values.get("unit_price"), values.get("currency"), row, "unit_price", errors)
                pack_size = values.get("pack_size", "")
                pack_unit = values.get("pack_unit", "")
                if bool(pack_size) != bool(pack_unit):
                    errors.append({**row, "field": "pack_size", "code": "PACK_SIZE_INCOMPLETE", "message": "Provide both pack_size and pack_unit, or leave both blank."})
                record["pack_size"] = _decimal_text(_parse_decimal(pack_size, row, "pack_size", errors)) if pack_size else None
                record["pack_unit"] = _unit(pack_unit, row, "pack_unit", errors) if pack_unit else None
                record["aliases"] = sorted({value.strip() for value in str(values.get("aliases", "")).split(";") if value.strip()})
                for field, maximum in (("sku", 64), ("name", 255), ("unit", 32)):
                    if isinstance(record.get(field), str) and len(record[field]) > maximum:
                        errors.append({**row, "field": field, "code": "FIELD_TOO_LONG", "message": f"{field} may contain at most {maximum} characters."})
                for alias in record["aliases"]:
                    if len(alias) > 255:
                        errors.append({**row, "field": "aliases", "code": "FIELD_TOO_LONG", "message": "Each alias may contain at most 255 characters."})
            elif kind == "inventory":
                record["sku"] = str(values.get("sku", "")).strip().upper()
                if not record["sku"]:
                    errors.append({**row, "field": "sku", "code": "REQUIRED_FIELD_MISSING", "message": "sku is required."})
                record["quantity_on_hand"] = _decimal_text(_parse_decimal(values.get("quantity_on_hand"), row, "quantity_on_hand", errors, allow_zero=True))
                record["unit"] = _unit(values.get("unit"), row, "unit", errors)
                if len(record["sku"]) > 64:
                    errors.append({**row, "field": "sku", "code": "FIELD_TOO_LONG", "message": "sku may contain at most 64 characters."})
            elif kind == "prices":
                record["sku"] = str(values.get("sku", "")).strip().upper()
                record["customer_id"] = str(values.get("customer_id", "")).strip() or None
                if not record["sku"]:
                    errors.append({**row, "field": "sku", "code": "REQUIRED_FIELD_MISSING", "message": "sku is required."})
                if len(record["sku"]) > 64 or len(record["customer_id"] or "") > 64:
                    errors.append({**row, "field": "sku", "code": "FIELD_TOO_LONG", "message": "sku and customer_id may contain at most 64 characters."})
                record["unit_price_minor"], record["currency"] = _money_minor(values.get("unit_price"), values.get("currency"), row, "unit_price", errors)
            else:
                record.update(
                    order_id=str(values.get("order_id", "")).strip(),
                    customer_id=str(values.get("customer_id", "")).strip(),
                    customer_name=str(values.get("customer_name", "")).strip(),
                    sku=str(values.get("sku", "")).strip().upper(),
                    unit=str(values.get("unit", "")).strip().casefold(),
                    status=str(values.get("status", "")).strip(),
                )
                if not record["order_id"] or (not record["customer_id"] and not record["customer_name"]) or not record["sku"] or not record["status"]:
                    errors.append({**row, "field": None, "code": "REQUIRED_FIELD_MISSING", "message": "order_id, customer_id or customer_name, sku, and status are required."})
                for field, maximum in (("order_id", 128), ("customer_id", 64), ("customer_name", 255), ("sku", 64), ("unit", 32), ("status", 32)):
                    if len(record[field]) > maximum:
                        errors.append({**row, "field": field, "code": "FIELD_TOO_LONG", "message": f"{field} may contain at most {maximum} characters."})
                record["quantity"] = _decimal_text(_parse_decimal(values.get("quantity") or values.get("quantity_on_hand"), row, "quantity", errors))
                record["unit"] = _unit(record["unit"], row, "unit", errors)
                record["unit_price_minor"], record["currency"] = _money_minor(values.get("unit_price"), values.get("currency"), row, "unit_price", errors)
                try:
                    parsed_date = datetime.fromisoformat(str(values.get("order_date", "")).strip().replace("Z", "+00:00"))
                    if parsed_date.tzinfo:
                        parsed_date = parsed_date.astimezone(timezone.utc).replace(tzinfo=None)
                    record["order_date"] = parsed_date.isoformat(sep=" ")
                except ValueError:
                    record["order_date"] = None
                    errors.append({**row, "field": "order_date", "code": "DATE_INVALID", "message": "order_date must be a valid ISO date or datetime."})
            records.append(record)
        normalized.append({"type": kind, "records": records})
    return normalized, errors


def _row_ref(record):
    return {"sheet": record["sheet"], "row": record["row"]}


def _record_key(record):
    return (record["sheet"], record["row"])


def _existing_snapshot(row, fields):
    if not row:
        return None
    result = {}
    for field in fields:
        value = row.get(field)
        result[field] = format(value, "f") if isinstance(value, Decimal) else int(value) if field.endswith("_minor") or field == "credit_used_minor" or field == "credit_limit_minor" else value
    return result


def _validation_error(errors, record, field, code, message):
    errors.append({**_row_ref(record), "field": field, "code": code, "message": message})


def build_import_plan(cursor, datasets, allow_updates, initial_errors=None):
    """Validate identifiers, relationships, units and conflicts against current DB state."""
    errors = list(initial_errors or [])
    plans = {key: [] for key in ("customers", "products", "inventory", "prices", "aliases", "previous_orders")}
    flat = [record for dataset in datasets for record in dataset["records"]]
    cursor.execute("SELECT id, name, credit_limit_minor, credit_used_minor, credit_currency FROM customers")
    db_customers = {row["id"]: row for row in cursor.fetchall()}
    db_customer_names = {}
    for row in db_customers.values():
        db_customer_names.setdefault(row["name"].strip().casefold(), []).append(row)
    cursor.execute("SELECT id, sku, name, quantity_on_hand, unit, unit_price_minor, price_currency, pack_size, pack_unit FROM products")
    db_products = {row["sku"].strip().upper(): row for row in cursor.fetchall()}
    cursor.execute("SELECT product_id, alias FROM product_aliases")
    db_aliases = {row["alias"].strip().casefold(): row["product_id"] for row in cursor.fetchall()}
    cursor.execute("SELECT customer_id, product_id, unit_price_minor, currency FROM customer_product_prices")
    db_prices = {(row["customer_id"], row["product_id"]): row for row in cursor.fetchall()}
    cursor.execute("SELECT customer_id, external_order_id FROM imported_order_history")
    db_history = {(row["customer_id"], row["external_order_id"]): True for row in cursor.fetchall()}

    bad_rows = {_record_key(error) for error in errors if "sheet" in error and "row" in error and error["row"] is not None}
    duplicate_seen = {}
    staged_aliases = {}
    for record in flat:
        kind = record["type"]
        values = {key: value for key, value in record.items() if key not in {"type", "sheet", "row"}}
        if kind == "customers":
            identifier = values.get("id")
            unique_key = ("customer", identifier.casefold()) if identifier else None
        elif kind in {"products", "inventory"}:
            identifier = values.get("sku")
            unique_key = ("product", identifier.casefold()) if identifier else None
        elif kind == "prices":
            identifier = (values.get("customer_id") or "catalog", values.get("sku"))
            unique_key = ("price", str(identifier).casefold()) if values.get("sku") else None
        elif kind == "previous_orders":
            identifier = (values.get("customer_id") or values.get("customer_name"), values.get("order_id"), values.get("sku"))
            unique_key = ("previous_order_line", str(identifier).casefold()) if values.get("order_id") and values.get("sku") else None
        else:
            unique_key = None
        if unique_key:
            if unique_key in duplicate_seen:
                _validation_error(errors, record, "identifier", "DUPLICATE_IDENTIFIER", "This identifier appears more than once in the uploaded file.")
                bad_rows.add(_record_key(record))
                bad_rows.add(_record_key(duplicate_seen[unique_key]))
            else:
                duplicate_seen[unique_key] = record

    staged_customers = {}
    staged_customer_names = {}
    staged_products = {}
    for record in flat:
        if _record_key(record) in bad_rows:
            continue
        if record["type"] == "customers":
            staged_customers[record["id"]] = record
            staged_customer_names.setdefault(record["name"].casefold(), []).append(record)
        elif record["type"] == "products":
            staged_products[record["sku"]] = record

    for name, records in staged_customer_names.items():
        if len(records) > 1:
            for record in records:
                _validation_error(errors, record, "name", "DUPLICATE_CUSTOMER_NAME", "Customer names must be unique in an import so future lookups are unambiguous.")
                bad_rows.add(_record_key(record))

    customer_plan_by_id = {}
    for record in flat:
        if record["type"] != "customers" or _record_key(record) in bad_rows:
            continue
        existing = db_customers.get(record["id"])
        competing = [row for row in db_customer_names.get(record["name"].casefold(), []) if row["id"] != record["id"]]
        if competing:
            _validation_error(errors, record, "name", "CUSTOMER_NAME_CONFLICT", "Another customer already has this name; resolve the duplicate before importing.")
            continue
        before = _existing_snapshot(existing, ("id", "name", "credit_limit_minor", "credit_used_minor", "credit_currency"))
        action = "update" if existing else "insert"
        if existing and not allow_updates:
            _validation_error(errors, record, "id", "EXISTING_RECORD_REQUIRES_UPDATE_CONFIRMATION", "This customer ID already exists. Preview again with updateExisting=true to review its proposed changes.")
        plan = {**_row_ref(record), "action": action if (not existing or allow_updates) else "blocked", "before": before, "values": {
            "id": record["id"], "name": record["name"], "credit_limit_minor": record["credit_limit_minor"],
            "credit_used_minor": record["credit_used_minor"], "credit_currency": record["currency"],
        }}
        plans["customers"].append(plan)
        customer_plan_by_id[record["id"]] = plan

    product_plan_by_sku = {}
    for record in flat:
        if record["type"] != "products" or _record_key(record) in bad_rows:
            continue
        existing = db_products.get(record["sku"])
        before = _existing_snapshot(existing, ("id", "sku", "name", "quantity_on_hand", "unit", "unit_price_minor", "price_currency", "pack_size", "pack_unit"))
        action = "update" if existing else "insert"
        if existing and not allow_updates:
            _validation_error(errors, record, "sku", "EXISTING_RECORD_REQUIRES_UPDATE_CONFIRMATION", "This SKU already exists. Preview again with updateExisting=true to review its proposed changes.")
        product_id = existing["id"] if existing else str(uuid.uuid4())
        plan = {**_row_ref(record), "action": action if (not existing or allow_updates) else "blocked", "before": before, "values": {
            "id": product_id, "sku": record["sku"], "name": record["name"],
            "quantity_on_hand": record["quantity_on_hand"], "unit": record["unit"],
            "unit_price_minor": record["unit_price_minor"], "price_currency": record["currency"],
            "pack_size": record["pack_size"], "pack_unit": record["pack_unit"],
        }}
        plans["products"].append(plan)
        product_plan_by_sku[record["sku"]] = plan

    # Do not allow two sheets to write the same product field in one import.
    product_skus = set(product_plan_by_sku)
    for record in flat:
        if record["type"] in {"inventory", "prices"} and record.get("sku") in product_skus:
            _validation_error(errors, record, "sku", "DUPLICATE_PRODUCT_UPDATE", "This import contains multiple datasets that change the same product; separate the imports to keep the preview unambiguous.")

    for record in flat:
        if _record_key(record) in bad_rows:
            continue
        kind = record["type"]
        if kind == "inventory":
            product_plan = product_plan_by_sku.get(record["sku"])
            existing = db_products.get(record["sku"])
            product = existing or (product_plan["values"] if product_plan else None)
            if not product:
                _validation_error(errors, record, "sku", "PRODUCT_NOT_FOUND", "Inventory rows must reference a product already in MySQL or in this same import.")
                continue
            if product["unit"].casefold() != record["unit"].casefold():
                _validation_error(errors, record, "unit", "UNIT_MISMATCH", f"The product is stored in {product['unit']}; inventory row uses {record['unit']}.")
                continue
            if product_plan:
                _validation_error(errors, record, "sku", "DUPLICATE_PRODUCT_UPDATE", "This product's stock is already set by the products sheet.")
                continue
            before = {"quantity_on_hand": format(Decimal(existing["quantity_on_hand"]), "f")}
            if not allow_updates:
                _validation_error(errors, record, "sku", "EXISTING_RECORD_REQUIRES_UPDATE_CONFIRMATION", "Updating existing inventory requires updateExisting=true and explicit confirmation.")
            plans["inventory"].append({**_row_ref(record), "action": "update" if allow_updates else "blocked", "before": before, "values": {"product_id": existing["id"], "sku": record["sku"], "quantity_on_hand": record["quantity_on_hand"]}})
        elif kind == "prices":
            product_plan = product_plan_by_sku.get(record["sku"])
            product = db_products.get(record["sku"]) or (product_plan["values"] if product_plan else None)
            if not product:
                _validation_error(errors, record, "sku", "PRODUCT_NOT_FOUND", "Price rows must reference a product already in MySQL or in this same import.")
            relationship_invalid = not product
            customer_id = record["customer_id"]
            if customer_id and customer_id not in db_customers and customer_id not in staged_customers:
                _validation_error(errors, record, "customer_id", "CUSTOMER_NOT_FOUND", "Customer-specific prices must reference an existing or same-import customer ID.")
                relationship_invalid = True
            if customer_id:
                customer = db_customers.get(customer_id) or staged_customers.get(customer_id)
                expected_currency = (customer.get("credit_currency") or customer.get("currency")) if customer else None
                if expected_currency and expected_currency != record["currency"]:
                    _validation_error(errors, record, "currency", "CUSTOMER_CURRENCY_MISMATCH", "Customer-specific price currency must match the customer's credit currency.")
                    relationship_invalid = True
            if relationship_invalid:
                continue
            if not customer_id:
                if product_plan:
                    _validation_error(errors, record, "sku", "DUPLICATE_PRODUCT_UPDATE", "This product's catalog price is already set by the products sheet.")
                    continue
                before = {"unit_price_minor": int(product["unit_price_minor"]), "price_currency": product["price_currency"]}
                action = "update"
                if not allow_updates:
                    _validation_error(errors, record, "sku", "EXISTING_RECORD_REQUIRES_UPDATE_CONFIRMATION", "Updating an existing catalog price requires updateExisting=true and explicit confirmation.")
                    action = "blocked"
                values = {"product_id": product["id"], "sku": record["sku"], "unit_price_minor": record["unit_price_minor"], "currency": record["currency"]}
            else:
                price_key = (customer_id, product["id"])
                existing_price = db_prices.get(price_key)
                before = {"unit_price_minor": int(existing_price["unit_price_minor"]), "currency": existing_price["currency"]} if existing_price else None
                action = "update" if existing_price else "insert"
                if existing_price and not allow_updates:
                    _validation_error(errors, record, "unit_price", "EXISTING_RECORD_REQUIRES_UPDATE_CONFIRMATION", "This customer price exists. Preview again with updateExisting=true to review its proposed change.")
                    action = "blocked"
                values = {"customer_id": customer_id, "product_id": product["id"], "sku": record["sku"], "unit_price_minor": record["unit_price_minor"], "currency": record["currency"]}
            plans["prices"].append({**_row_ref(record), "action": action, "before": before, "values": values})
        elif kind == "products":
            product_plan = product_plan_by_sku.get(record["sku"])
            if not product_plan:
                continue
            for alias in record["aliases"]:
                owner = db_aliases.get(alias.casefold())
                staged_owner = staged_aliases.get(alias.casefold())
                collision = any(row["sku"].casefold() == alias.casefold() or row["name"].casefold() == alias.casefold() for row in db_products.values() if row["id"] != product_plan["values"]["id"])
                if (owner and owner != product_plan["values"]["id"]) or (staged_owner and staged_owner != product_plan["values"]["id"]) or collision:
                    _validation_error(errors, record, "aliases", "ALIAS_CONFLICT", f"Alias '{alias}' is already associated with another product.")
                    continue
                staged_aliases[alias.casefold()] = product_plan["values"]["id"]
                plans["aliases"].append({**_row_ref(record), "action": "skip" if owner == product_plan["values"]["id"] else "insert", "values": {"product_id": product_plan["values"]["id"], "alias": alias}})

    # Historical orders are kept outside the operational orders table and never adjust inventory or credit.
    history_groups = {}
    for record in flat:
        if record["type"] != "previous_orders" or _record_key(record) in bad_rows:
            continue
        customer_id = record["customer_id"]
        if not customer_id and record["customer_name"]:
            matches = db_customer_names.get(record["customer_name"].casefold(), [])
            staged_matches = staged_customer_names.get(record["customer_name"].casefold(), [])
            if len(matches) + len(staged_matches) == 1:
                customer_id = matches[0]["id"] if matches else staged_matches[0]["id"]
            else:
                _validation_error(errors, record, "customer_name", "CUSTOMER_RELATIONSHIP_AMBIGUOUS", "Historical order customer must resolve to exactly one existing or same-import customer.")
                continue
        customer = db_customers.get(customer_id) or customer_plan_by_id.get(customer_id, {}).get("values")
        if customer and record["customer_name"] and customer["name"].strip().casefold() != record["customer_name"].strip().casefold():
            _validation_error(errors, record, "customer_name", "CUSTOMER_ID_NAME_MISMATCH", "The supplied customer ID and name refer to different customer records.")
            continue
        product_plan = product_plan_by_sku.get(record["sku"])
        product = db_products.get(record["sku"]) or (product_plan["values"] if product_plan else None)
        if not customer:
            _validation_error(errors, record, "customer_id", "CUSTOMER_NOT_FOUND", "Historical order must reference an existing or same-import customer.")
            continue
        if not product:
            _validation_error(errors, record, "sku", "PRODUCT_NOT_FOUND", "Historical order must reference an existing or same-import product.")
            continue
        if product["unit"].casefold() != record["unit"]:
            _validation_error(errors, record, "unit", "UNIT_MISMATCH", f"The product is stored in {product['unit']}; historical row uses {record['unit']}.")
            continue
        group_key = (customer["id"], record["order_id"])
        line = {
            "product_id": product["id"], "sku": record["sku"], "quantity": record["quantity"], "unit": record["unit"],
            "unit_price_minor": record["unit_price_minor"], "currency": record["currency"],
        }
        group = history_groups.setdefault(group_key, {
            "source": _row_ref(record), "customer_id": customer["id"], "external_order_id": record["order_id"],
            "order_date": record["order_date"], "status": record["status"], "currency": record["currency"], "lines": [],
        })
        if group["order_date"] != record["order_date"] or group["currency"] != record["currency"] or group["status"].casefold() != record["status"].casefold():
            _validation_error(errors, record, None, "HISTORICAL_ORDER_INCONSISTENT", "Rows for one previous order must have the same date, status, and currency.")
            continue
        if any(old["sku"] == line["sku"] for old in group["lines"]):
            _validation_error(errors, record, "sku", "DUPLICATE_ORDER_LINE", "This order contains the same SKU more than once; consolidate it into one row.")
            continue
        group["lines"].append(line)

    for (customer_id, external_id), group in history_groups.items():
        if (customer_id, external_id) in db_history:
            _validation_error(errors, {"sheet": group["source"]["sheet"], "row": group["source"]["row"]}, "order_id", "HISTORICAL_ORDER_EXISTS", "This previous order is already in imported history and will not be overwritten.")
            continue
        for line in group["lines"]:
            line["line_total_minor"] = int((Decimal(line["quantity"]) * Decimal(line["unit_price_minor"])).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        total = sum(line["line_total_minor"] for line in group["lines"])
        plans["previous_orders"].append({
            **group["source"], "action": "insert", "before": None,
            "values": {**{key: group[key] for key in ("customer_id", "external_order_id", "order_date", "status", "currency")}, "total_minor": total, "lines": group["lines"]},
        })

    summary = {
        "rows": len(flat),
        "errors": len(errors),
        "creates": sum(row["action"] == "insert" for rows in plans.values() for row in rows),
        "updates": sum(row["action"] == "update" for rows in plans.values() for row in rows),
        "skipped": sum(row["action"] == "skip" for rows in plans.values() for row in rows),
        "byType": {kind: len(rows) for kind, rows in plans.items()},
    }
    return {"plans": plans, "errors": errors, "summary": summary}
